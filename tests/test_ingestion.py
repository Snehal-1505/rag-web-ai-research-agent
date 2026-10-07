"""
Unit tests for Document Loader (Phase 2 Ingestion)
"""
import pytest
from pathlib import Path
import fitz  # PyMuPDF
import docx

from src.ingestion.document_loader import (
    load_document,
    load_multiple_documents,
    load_txt,
    load_pdf,
    load_docx,
)


@pytest.fixture
def temp_sample_files(tmp_path):
    """Fixture to create sample PDF, TXT, and DOCX files for testing."""
    # 1. Create sample TXT file
    txt_path = tmp_path / "sample.txt"
    txt_path.write_text("This is a sample text file for testing document ingestion.", encoding="utf-8")

    # 2. Create sample PDF file
    pdf_path = tmp_path / "sample.pdf"
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((50, 50), "This is page 1 of sample PDF.")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "This is page 2 of sample PDF.")
    pdf_doc.save(pdf_path)
    pdf_doc.close()

    # 3. Create sample DOCX file
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Sample Document", level=1)
    doc.add_paragraph("This is a sample DOCX file paragraph.")
    doc.save(docx_path)

    return {
        "txt": txt_path,
        "pdf": pdf_path,
        "docx": docx_path,
    }


def test_load_txt(temp_sample_files):
    txt_file = temp_sample_files["txt"]
    docs = load_document(txt_file)
    assert len(docs) == 1
    assert "sample text file" in docs[0].content
    assert docs[0].meta["filename"] == "sample.txt"
    assert docs[0].meta["file_type"] == "txt"
    assert docs[0].meta["page_number"] == 1
    assert "upload_time" in docs[0].meta


def test_load_pdf(temp_sample_files):
    pdf_file = temp_sample_files["pdf"]
    docs = load_document(pdf_file)
    assert len(docs) == 2  # 2 pages
    assert "page 1" in docs[0].content
    assert docs[0].meta["page_number"] == 1
    assert "page 2" in docs[1].content
    assert docs[1].meta["page_number"] == 2
    assert docs[0].meta["file_type"] == "pdf"


def test_load_docx(temp_sample_files):
    docx_file = temp_sample_files["docx"]
    docs = load_document(docx_file)
    assert len(docs) == 1
    assert "sample DOCX file" in docs[0].content
    assert docs[0].meta["file_type"] == "docx"
    assert docs[0].meta["filename"] == "sample.docx"


def test_load_multiple_documents(temp_sample_files):
    file_list = [temp_sample_files["txt"], temp_sample_files["pdf"], temp_sample_files["docx"]]
    docs = load_multiple_documents(file_list)
    assert len(docs) == 4  # 1 txt + 2 pdf pages + 1 docx


def test_unsupported_file_format(tmp_path):
    invalid_file = tmp_path / "sample.xyz"
    invalid_file.write_text("Test content", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported file format"):
        load_document(invalid_file)
