"""
Document Loader Module for Haystack AI RAG Application.

Supports loading PDF, TXT, and DOCX files into Haystack Document objects while
preserving metadata (filename, file_type, page_number, source, upload_time).
"""
import os
from datetime import datetime
from pathlib import Path
from typing import List, Union

try:
    import pymupdf as fitz  # PyMuPDF modern import name
except ImportError:
    import fitz  # PyMuPDF legacy import name
import docx  # python-docx for DOCX extraction
from haystack import Document


def get_current_timestamp() -> str:
    """Returns ISO format timestamp for metadata tracking."""
    return datetime.now().isoformat()


def load_txt(file_path: Path) -> List[Document]:
    """
    Loads a TXT file and returns a list containing one Haystack Document.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            text = f.read()

    text = text.strip()
    if not text:
        return []

    metadata = {
        "filename": file_path.name,
        "file_type": "txt",
        "page_number": 1,
        "source": str(file_path),
        "upload_time": get_current_timestamp(),
    }

    return [Document(content=text, meta=metadata)]


def load_pdf(file_path: Path) -> List[Document]:
    """
    Loads a PDF file using PyMuPDF (fitz) and creates a Haystack Document per page.
    """
    documents: List[Document] = []
    pdf_doc = fitz.open(file_path)

    timestamp = get_current_timestamp()

    for page_num in range(len(pdf_doc)):
        page = pdf_doc.load_page(page_num)
        text = page.get_text("text").strip()

        if text:  # Skip empty pages
            metadata = {
                "filename": file_path.name,
                "file_type": "pdf",
                "page_number": page_num + 1,
                "source": str(file_path),
                "upload_time": timestamp,
            }
            documents.append(Document(content=text, meta=metadata))

    pdf_doc.close()
    return documents


def load_docx(file_path: Path) -> List[Document]:
    """
    Loads a DOCX file using python-docx and creates a Haystack Document.
    """
    doc = docx.Document(file_path)
    full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

    if not full_text.strip():
        return []

    metadata = {
        "filename": file_path.name,
        "file_type": "docx",
        "page_number": 1,
        "source": str(file_path),
        "upload_time": get_current_timestamp(),
    }

    return [Document(content=full_text, meta=metadata)]


def load_document(file_path: Union[str, Path]) -> List[Document]:
    """
    Main document loading router. Determines file type by extension and delegates
    to the appropriate loader function.

    Supported formats: .pdf, .txt, .docx

    Returns:
        List[Document]: A list of Haystack Document objects with metadata attached.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()

    if ext == ".txt":
        return load_txt(path)
    elif ext == ".pdf":
        return load_pdf(path)
    elif ext == ".docx":
        return load_docx(path)
    else:
        raise ValueError(
            f"Unsupported file format '{ext}'. Only .pdf, .txt, and .docx are supported."
        )


def load_multiple_documents(file_paths: List[Union[str, Path]]) -> List[Document]:
    """
    Helper function to load multiple files and aggregate their Haystack Documents.
    """
    all_docs: List[Document] = []
    for fp in file_paths:
        docs = load_document(fp)
        all_docs.extend(docs)
    return all_docs
